# tests/integration/test_real_service_extraction.py
"""
End-to-end extraction through the *real* services.

Every other handler test injects mocked services, which means a handler can
stop emitting image tags or page markers entirely and the suite stays green.
These tests build small real documents, run them through a real
``DocumentProcessor``, and assert on the produced text — the only layer that
notices when the image or tagging pipeline quietly stops working.
"""

from __future__ import annotations

import base64
import io
from pathlib import Path

import pytest

from contextifier import DocumentProcessor, ProcessingConfig

# 1×1 red PNG — smallest thing every image library accepts.
_PNG_1x1 = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg=="
)


@pytest.fixture()
def processor(tmp_path: Path) -> DocumentProcessor:
    """Processor whose images land in a temp dir instead of ``temp/images``."""
    config = ProcessingConfig().with_images(directory_path=str(tmp_path / "images"))
    return DocumentProcessor(config=config)


# ── DOCX ──────────────────────────────────────────────────────────────────

def _build_docx(path: Path) -> None:
    docx = pytest.importorskip("docx")
    from docx.shared import Inches

    doc = docx.Document()
    doc.add_paragraph("First page body text.")
    doc.add_picture(io.BytesIO(_PNG_1x1), width=Inches(1))
    doc.add_page_break()
    doc.add_paragraph("Second page body text.")
    doc.save(path)


def test_docx_emits_image_tag_and_page_markers(
    processor: DocumentProcessor, tmp_path: Path
) -> None:
    path = tmp_path / "sample.docx"
    _build_docx(path)

    text = processor.extract_text(str(path))

    assert "[Image:" in text, "inline DOCX image produced no image tag"
    assert "[Page Number: 1]" in text
    assert "[Page Number: 2]" in text, "page break did not open a new page marker"
    assert "Second page body text." in text


def test_docx_chunks_carry_page_numbers(
    processor: DocumentProcessor, tmp_path: Path
) -> None:
    path = tmp_path / "sample.docx"
    _build_docx(path)

    result = processor.extract_chunks(
        str(path), chunk_size=200, include_position_metadata=True
    )

    assert result.chunks_with_metadata
    pages = {c.metadata.page_number for c in result.chunks_with_metadata}
    assert pages != {None}, "no chunk carries a page number"


# ── PPTX ──────────────────────────────────────────────────────────────────

def _build_pptx(path: Path) -> None:
    pptx = pytest.importorskip("pptx")
    from pptx.util import Inches

    prs = pptx.Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Slide One Title"
    slide.shapes.add_picture(io.BytesIO(_PNG_1x1), Inches(1), Inches(1), Inches(1))
    second = prs.slides.add_slide(prs.slide_layouts[5])
    second.shapes.title.text = "Slide Two Title"
    prs.save(path)


def test_pptx_emits_image_tag_and_slide_markers(
    processor: DocumentProcessor, tmp_path: Path
) -> None:
    path = tmp_path / "deck.pptx"
    _build_pptx(path)

    text = processor.extract_text(str(path))

    assert "[Image:" in text, "PPTX picture produced no image tag"
    assert "[Slide Number: 1]" in text
    assert "[Slide Number: 2]" in text
    assert "Slide Two Title" in text


# ── XLSX ──────────────────────────────────────────────────────────────────

def _build_xlsx(path: Path) -> None:
    openpyxl = pytest.importorskip("openpyxl")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["Name", "Score"])
    ws.append(["Alice", 95])
    ws.append(["Bob", 88])
    wb.save(path)


def test_xlsx_emits_sheet_marker(processor: DocumentProcessor, tmp_path: Path) -> None:
    path = tmp_path / "book.xlsx"
    _build_xlsx(path)

    text = processor.extract_text(str(path))

    assert "[Sheet: Data]" in text
    assert "Alice" in text


# ── Tag configuration must reach the handlers ────────────────────────────

def test_custom_page_tag_configuration_is_honoured(tmp_path: Path) -> None:
    """A caller that reconfigures the tag format must see it in the output.

    The handlers used to hardcode their own format, so this silently did
    nothing — and the page-aware chunking strategy, which looks for the
    *configured* prefix, never matched.
    """
    docx = pytest.importorskip("docx")

    path = tmp_path / "tagged.docx"
    doc = docx.Document()
    doc.add_paragraph("Body")
    doc.save(path)

    config = ProcessingConfig().with_tags(page_prefix="<page>", page_suffix="</page>")
    text = DocumentProcessor(config=config).extract_text(str(path))

    assert "<page>1</page>" in text
    assert "[Page Number:" not in text


# ── HWPX ──────────────────────────────────────────────────────────────────

_HWPX_NS = (
    'xmlns:hp="http://www.hancom.co.kr/hwpml/2011/paragraph" '
    'xmlns:hc="http://www.hancom.co.kr/hwpml/2011/core" '
    'xmlns:hs="http://www.hancom.co.kr/hwpml/2011/section"'
)


def _build_hwpx(path: Path) -> None:
    """Minimal but structurally faithful HWPX: body, table, shape, header."""
    import zipfile

    def para(text: str) -> str:
        return f"<hp:p><hp:run><hp:t>{text}</hp:t></hp:run></hp:p>"

    def cell(row: int, col: int, text: str) -> str:
        return (
            f'<hp:tc><hp:cellAddr colAddr="{col}" rowAddr="{row}"/>'
            f'<hp:cellSpan colSpan="1" rowSpan="1"/>'
            f"<hp:subList>{para(text)}</hp:subList></hp:tc>"
        )

    table = (
        '<hp:tbl rowCnt="2" colCnt="2">'
        f'<hp:tr>{cell(0, 0, "R1C1")}{cell(0, 1, "R1C2")}</hp:tr>'
        f'<hp:tr>{cell(1, 0, "R2C1")}{cell(1, 1, "R2C2")}</hp:tr>'
        "</hp:tbl>"
    )
    section = (
        f"<hs:sec {_HWPX_NS}>"
        + para("Body paragraph one.")
        + f"<hp:p><hp:run>{table}</hp:run></hp:p>"
        + "<hp:p><hp:run><hp:rect><hp:drawText><hp:subList>"
        + para("Diagram caption text")
        + "</hp:subList></hp:drawText></hp:rect></hp:run></hp:p>"
        + "<hp:p><hp:run><hp:ctrl><hp:header><hp:subList>"
        + para("Running page header")
        + "</hp:subList></hp:header></hp:ctrl></hp:run></hp:p>"
        + "</hs:sec>"
    )

    with zipfile.ZipFile(path, "w") as zf:
        zf.writestr("mimetype", "application/hwp+zip")
        zf.writestr("version.xml", '<?xml version="1.0"?><hv:HCFVersion/>')
        zf.writestr("Contents/section0.xml", section)


def test_hwpx_structure_is_not_duplicated(
    processor: DocumentProcessor, tmp_path: Path
) -> None:
    path = tmp_path / "doc.hwpx"
    _build_hwpx(path)

    text = processor.extract_text(str(path))

    assert "<table>" in text
    for cell_text in ("R1C1", "R1C2", "R2C1", "R2C2"):
        assert text.count(cell_text) == 1, (
            f"{cell_text} appears {text.count(cell_text)}× — cell paragraphs are "
            "being emitted a second time as body text"
        )


def test_hwpx_shape_text_and_header_are_placed_correctly(
    processor: DocumentProcessor, tmp_path: Path
) -> None:
    path = tmp_path / "doc.hwpx"
    _build_hwpx(path)

    text = processor.extract_text(str(path))

    assert "Diagram caption text" in text
    assert text.index("Body paragraph one.") < text.index("Diagram caption text")

    assert "[Headers]" in text
    assert "Running page header" in text
    assert text.index("Diagram caption text") < text.index("[Headers]")
