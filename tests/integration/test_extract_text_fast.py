# tests/integration/test_extract_text_fast.py
"""
The cheap path: words only.

`extract_text_fast` answers "does this file contain a forbidden word or a
piece of personal data?". It gives up structure, images, charts, OCR and
metadata in exchange for skipping the analysis that makes the full pipeline
slow — and it must still return every word a scan would need to find.
"""

from __future__ import annotations

import base64
import io
import time
from pathlib import Path

import pytest

from contextifier import DocumentProcessor

_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg=="
)

SECRET = "주민등록번호 900101-1234567"


@pytest.fixture()
def processor() -> DocumentProcessor:
    return DocumentProcessor()


class TestDocx:
    @pytest.fixture()
    def path(self, tmp_path: Path) -> Path:
        docx = pytest.importorskip("docx")
        from docx.shared import Inches

        d = docx.Document()
        d.add_paragraph("Ordinary paragraph.")
        d.add_picture(io.BytesIO(_PNG), width=Inches(1))
        table = d.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Field"
        table.cell(0, 1).text = SECRET
        out = tmp_path / "scan.docx"
        d.save(out)
        return out

    def test_finds_text_in_cells(self, processor: DocumentProcessor, path: Path) -> None:
        assert SECRET in processor.extract_text_fast(str(path))

    def test_omits_structure_and_metadata(
        self, processor: DocumentProcessor, path: Path
    ) -> None:
        text = processor.extract_text_fast(str(path))
        assert "[Document-Metadata]" not in text
        assert "<table>" not in text
        assert "[Image:" not in text


class TestXlsx:
    def test_finds_cell_values(self, processor: DocumentProcessor, tmp_path: Path) -> None:
        openpyxl = pytest.importorskip("openpyxl")

        path = tmp_path / "scan.xlsx"
        wb = openpyxl.Workbook()
        wb.active.append(["name", "id"])
        wb.active.append(["Alice", SECRET])
        wb.save(path)

        assert SECRET in processor.extract_text_fast(str(path))


class TestPptx:
    def test_finds_slide_text(self, processor: DocumentProcessor, tmp_path: Path) -> None:
        pptx = pytest.importorskip("pptx")

        path = tmp_path / "scan.pptx"
        prs = pptx.Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[5])
        slide.shapes.title.text = SECRET
        prs.save(path)

        assert SECRET in processor.extract_text_fast(str(path))


class TestPdf:
    @pytest.fixture()
    def path(self) -> Path:
        full = Path(__file__).resolve().parents[2] / "_deep_test/test_files/multipage.pdf"
        if not full.exists():
            pytest.skip("multipage.pdf not available")
        return full

    def test_returns_page_text(self, processor: DocumentProcessor, path: Path) -> None:
        text = processor.extract_text_fast(str(path))
        assert "Page 1: Chapter Title" in text
        assert "[Page Number:" not in text

    def test_is_faster_than_the_full_pipeline(
        self, processor: DocumentProcessor, path: Path
    ) -> None:
        start = time.perf_counter()
        processor.extract_text_fast(str(path))
        fast = time.perf_counter() - start

        start = time.perf_counter()
        processor.extract_text(str(path))
        full = time.perf_counter() - start

        assert fast < full, f"fast={fast:.3f}s full={full:.3f}s"


class TestImageFiles:
    def test_image_file_scans_as_empty(
        self, processor: DocumentProcessor, tmp_path: Path
    ) -> None:
        """OCR is the cost this path exists to avoid."""
        path = tmp_path / "photo.png"
        path.write_bytes(_PNG)
        assert processor.extract_text_fast(str(path)) == ""


class TestUnsupported:
    def test_unknown_extension_raises(
        self, processor: DocumentProcessor, tmp_path: Path
    ) -> None:
        from contextifier import UnsupportedFormatError

        path = tmp_path / "file.unknownext"
        path.write_text("x")
        with pytest.raises(UnsupportedFormatError):
            processor.extract_text_fast(str(path))


class TestPhrasesSurviveRunSplitting:
    def test_a_phrase_split_across_runs_is_findable(self, tmp_path: Path) -> None:
        """Word splits a phrase at every formatting change.

        Reading one line per text node would cut a searched-for phrase into
        pieces, and a scan for it would come back empty on a document that
        plainly contains it.
        """
        docx = pytest.importorskip("docx")

        d = docx.Document()
        para = d.add_paragraph()
        para.add_run("경희대학교 ")
        para.add_run("경영").bold = True
        para.add_run("연구원")
        path = tmp_path / "runs.docx"
        d.save(path)

        text = DocumentProcessor().extract_text_fast(str(path))
        assert "경희대학교 경영연구원" in text


class TestHeadersAndFooters:
    def test_header_text_is_scanned(self, tmp_path: Path) -> None:
        docx = pytest.importorskip("docx")

        d = docx.Document()
        d.sections[0].header.paragraphs[0].text = SECRET
        d.add_paragraph("body")
        path = tmp_path / "hdr.docx"
        d.save(path)

        assert SECRET in DocumentProcessor().extract_text_fast(str(path))
