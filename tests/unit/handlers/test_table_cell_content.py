# tests/unit/handlers/test_table_cell_content.py
"""
What a table cell can hold besides typed text.

A cell commonly carries a pasted screenshot instead of text, a content control
instead of a plain paragraph, or a sub-table. Reading cells for ``w:t`` alone
returned an empty cell in each case — and for the image that is not merely
unrendered but unreachable, because the OCR pass works from the tags left in
the text.
"""

from __future__ import annotations

import base64
import io
from pathlib import Path

import pytest

from contextifier import DocumentProcessor, ProcessingConfig

_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg=="
)


@pytest.fixture()
def processor(tmp_path: Path) -> DocumentProcessor:
    config = ProcessingConfig().with_images(directory_path=str(tmp_path / "images"))
    return DocumentProcessor(config=config)


class TestDocxCellImage:
    @pytest.fixture()
    def doc_with_cell_image(self, tmp_path: Path) -> Path:
        docx = pytest.importorskip("docx")
        from docx.shared import Inches

        d = docx.Document()
        table = d.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Label"
        table.cell(0, 1).text = "Value"
        table.cell(1, 0).text = "Screenshot"
        table.cell(1, 1).paragraphs[0].add_run().add_picture(
            io.BytesIO(_PNG), width=Inches(1)
        )
        path = tmp_path / "cell_image.docx"
        d.save(path)
        return path

    def test_image_tag_lands_in_its_own_cell(
        self, processor: DocumentProcessor, doc_with_cell_image: Path
    ) -> None:
        text = processor.extract_text(str(doc_with_cell_image))

        assert "[Image:" in text, "the cell image produced no tag at all"
        row = [
            line
            for line in text.splitlines()
            if "Screenshot" in line
        ]
        assert row, "the row holding the screenshot is missing"
        assert "[Image:" in row[0], (
            "the tag must sit in the cell that holds the image:\n" + row[0]
        )

    def test_caption_and_image_both_survive(
        self, processor: DocumentProcessor, tmp_path: Path
    ) -> None:
        docx = pytest.importorskip("docx")
        from docx.shared import Inches

        d = docx.Document()
        table = d.add_table(rows=1, cols=1)
        para = table.cell(0, 0).paragraphs[0]
        para.text = "Figure 3 caption"
        para.add_run().add_picture(io.BytesIO(_PNG), width=Inches(1))
        path = tmp_path / "caption.docx"
        d.save(path)

        text = processor.extract_text(str(path))
        assert "Figure 3 caption" in text
        assert "[Image:" in text


class TestDocxNestedTable:
    def test_sub_table_inside_a_cell_is_kept(
        self, processor: DocumentProcessor, tmp_path: Path
    ) -> None:
        docx = pytest.importorskip("docx")

        d = docx.Document()
        outer = d.add_table(rows=1, cols=2)
        outer.cell(0, 0).text = "outer-left"
        inner = outer.cell(0, 1).add_table(rows=1, cols=2)
        inner.cell(0, 0).text = "INNER-A"
        inner.cell(0, 1).text = "INNER-B"
        path = tmp_path / "nested.docx"
        d.save(path)

        text = processor.extract_text(str(path))
        assert "INNER-A" in text, "the nested table was dropped"
        assert "INNER-B" in text


class TestDocxCellContentControl:
    def test_control_inside_a_cell_is_unwrapped(
        self, processor: DocumentProcessor, tmp_path: Path
    ) -> None:
        docx = pytest.importorskip("docx")
        from docx.oxml.ns import qn

        d = docx.Document()
        table = d.add_table(rows=1, cols=1)
        cell = table.cell(0, 0)
        para = cell.paragraphs[0]
        sdt = para._p.makeelement(qn("w:sdt"), {})
        content = para._p.makeelement(qn("w:sdtContent"), {})
        inner_p = para._p.makeelement(qn("w:p"), {})
        run = para._p.makeelement(qn("w:r"), {})
        t = para._p.makeelement(qn("w:t"), {})
        t.text = "CELL-CONTROL-VALUE"
        run.append(t)
        inner_p.append(run)
        content.append(inner_p)
        sdt.append(content)
        cell._tc.append(sdt)
        path = tmp_path / "cell_sdt.docx"
        d.save(path)

        text = processor.extract_text(str(path))
        assert "CELL-CONTROL-VALUE" in text
