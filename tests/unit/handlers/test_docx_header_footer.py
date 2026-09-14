# tests/unit/handlers/test_docx_header_footer.py
"""
Header / footer extraction for DOCX.

Headers and footers are ordinary block containers: they hold paragraphs,
tables, content controls and shapes just like the body does. Reading only
``header.paragraphs`` silently drops the rest — and a table is the usual way
a document puts a document number, revision and classification in its header.
"""

from __future__ import annotations

from pathlib import Path

import pytest

from contextifier import DocumentProcessor


@pytest.fixture()
def doc_with_rich_header(tmp_path: Path) -> Path:
    docx = pytest.importorskip("docx")

    d = docx.Document()
    section = d.sections[0]
    section.header.paragraphs[0].text = "CONFIDENTIAL HEADER"
    table = section.header.add_table(rows=2, cols=2, width=section.page_width)
    table.cell(0, 0).text = "Doc No."
    table.cell(0, 1).text = "HDR-2026-001"
    table.cell(1, 0).text = "Revision"
    table.cell(1, 1).text = "HDR-REV-B"
    section.footer.paragraphs[0].text = "PAGE FOOTER TEXT"
    d.add_paragraph("body paragraph")

    path = tmp_path / "rich_header.docx"
    d.save(path)
    return path


def test_header_paragraph_and_footer_are_extracted(doc_with_rich_header: Path) -> None:
    text = DocumentProcessor().extract_text(str(doc_with_rich_header))
    assert "CONFIDENTIAL HEADER" in text
    assert "PAGE FOOTER TEXT" in text


def test_header_table_cells_are_extracted(doc_with_rich_header: Path) -> None:
    text = DocumentProcessor().extract_text(str(doc_with_rich_header))
    assert "HDR-2026-001" in text, "table inside the header was dropped"
    assert "HDR-REV-B" in text


def test_header_content_control_is_extracted(tmp_path: Path) -> None:
    """A content control in a header must be unwrapped like anywhere else."""
    docx = pytest.importorskip("docx")
    from docx.oxml.ns import qn

    d = docx.Document()
    header = d.sections[0].header
    para = header.paragraphs[0]
    para.text = "Title: "
    sdt = para._p.makeelement(qn("w:sdt"), {})
    content = para._p.makeelement(qn("w:sdtContent"), {})
    run = para._p.makeelement(qn("w:r"), {})
    t = para._p.makeelement(qn("w:t"), {})
    t.text = "HDR-CONTROL-VALUE"
    run.append(t)
    content.append(run)
    sdt.append(content)
    para._p.append(sdt)
    d.add_paragraph("body")

    path = tmp_path / "sdt_header.docx"
    d.save(path)

    text = DocumentProcessor().extract_text(str(path))
    assert "HDR-CONTROL-VALUE" in text
