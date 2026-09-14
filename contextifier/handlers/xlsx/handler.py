# contextifier/handlers/xlsx/handler.py
"""
XLSXHandler — Handler for modern Excel XLSX spreadsheets (.xlsx ONLY).

XLSX is an OOXML (Office Open XML) ZIP-based format parsed with openpyxl.
This is fundamentally different from legacy .xls (BIFF binary) which
requires xlrd or LibreOffice conversion.

Pipeline:
    Convert:  Raw bytes → openpyxl Workbook (data_only=True)
    Preprocess: Wrap Workbook, pre-extract charts/images/textboxes from ZIP
    Metadata: OOXML core properties → DocumentMetadata
    Content:  Per-sheet layout detection → table conversion (MD/HTML),
              chart extraction, image extraction, textbox extraction
    Postprocess: Assemble with sheet tags and metadata block
"""

from __future__ import annotations

from typing import Any, FrozenSet, Optional

from contextifier.handlers.base import BaseHandler
from contextifier.handlers.html._detect import looks_like_html
from contextifier.pipeline.converter import BaseConverter
from contextifier.pipeline.preprocessor import BasePreprocessor
from contextifier.pipeline.metadata_extractor import BaseMetadataExtractor
from contextifier.pipeline.content_extractor import BaseContentExtractor
from contextifier.pipeline.postprocessor import BasePostprocessor, DefaultPostprocessor

from contextifier.types import ExtractionResult, FileContext

from contextifier.handlers.xlsx.converter import XlsxConverter
from contextifier.handlers.xlsx.preprocessor import XlsxPreprocessor
from contextifier.handlers.xlsx.metadata_extractor import XlsxMetadataExtractor
from contextifier.handlers.xlsx.content_extractor import XlsxContentExtractor


class XLSXHandler(BaseHandler):
    """Handler for modern Excel files (.xlsx only)."""

    @property
    def supported_extensions(self) -> FrozenSet[str]:
        return frozenset({"xlsx"})

    @property
    def handler_name(self) -> str:
        return "XLSX Handler"

    # ── Delegation ───────────────────────────────────────────────────────

    def _check_delegation(
        self,
        file_context: FileContext,
        **kwargs: Any,
    ) -> Optional[ExtractionResult]:
        """Delegate when the .xlsx file is really an HTML export.

        Report portals serve HTML tables under a spreadsheet filename; openpyxl
        cannot open them, so without this the whole document fails to convert.
        """
        if looks_like_html(file_context.get("file_data", b"")):
            self._logger.info("XLSX file is actually an HTML export")
            return self._delegate_to(
                "html",
                file_context,
                include_metadata=kwargs.get("include_metadata", True),
                **{k: v for k, v in kwargs.items() if k != "include_metadata"},
            )
        return None

    # ── Pipeline stages ──────────────────────────────────────────────────

    def extract_text_fast(self, file_context: FileContext, **kwargs: Any) -> str:
        """Cell values only — no layout detection, images or charts."""
        import io

        try:
            from openpyxl import load_workbook

            workbook = load_workbook(
                io.BytesIO(file_context.get("file_data", b"")),
                read_only=True,
                data_only=True,
            )
        except Exception as exc:
            self._logger.debug("Fast XLSX path unavailable (%s); using pipeline", exc)
            return super().extract_text_fast(file_context, **kwargs)

        try:
            lines = []
            for sheet in workbook.worksheets:
                for row in sheet.iter_rows(values_only=True):
                    values = [
                        str(value).strip()
                        for value in row
                        if value is not None and str(value).strip()
                    ]
                    if values:
                        lines.append("\t".join(values))
            return "\n".join(lines)
        finally:
            workbook.close()

    def create_converter(self) -> BaseConverter:
        xlsx_opts = self._config.format_options.get("xlsx", {})
        data_only = xlsx_opts.get("data_only", True)
        read_only = xlsx_opts.get("read_only", False)
        return XlsxConverter(data_only=data_only, read_only=read_only)

    def create_preprocessor(self) -> BasePreprocessor:
        return XlsxPreprocessor()

    def create_metadata_extractor(self) -> BaseMetadataExtractor:
        return XlsxMetadataExtractor()

    def create_content_extractor(self) -> BaseContentExtractor:
        xlsx_opts = self._config.format_options.get("xlsx", {})
        include_hidden = xlsx_opts.get("include_hidden_sheets", False)
        return XlsxContentExtractor(
            image_service=self._image_service,
            tag_service=self._tag_service,
            chart_service=self._chart_service,
            table_service=self._table_service,
            include_hidden_sheets=include_hidden,
        )

    def create_postprocessor(self) -> BasePostprocessor:
        return DefaultPostprocessor(
            self._config,
            metadata_service=self._metadata_service,
            tag_service=self._tag_service,
        )


__all__ = ["XLSXHandler"]
